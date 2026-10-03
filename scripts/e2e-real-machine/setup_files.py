#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""真机 180 场景测试：种子文件生成。
产出（artifacts/e2e-180/files/）：
  doc-A-内容.docx … doc-G-脱敏.docx（7 个 Writer 组文档，含预埋内容）
  e2e-et.xlsx（60 个工作表 S01..S60，标准订单数据 8 列 × 21 行）
  e2e-wpp.pptx（6 页种子演示）
幂等：重复运行整体覆盖重建（文件未在 WPS 中打开为前提）。"""
import base64
import os
import random

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "..", "artifacts", "e2e-180", "files")
os.makedirs(OUT, exist_ok=True)

import struct
import zlib


def make_png(width=80, height=40, rgb=(58, 110, 240)):
    """生成纯色 PNG（无需 PIL），供 docx 嵌图与 picture_add 场景使用。"""
    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


PNG_PATH = os.path.join(OUT, "seed-image.png")
with open(PNG_PATH, "wb") as f:
    f.write(make_png())


def para(doc, text, style=None):
    p = doc.add_paragraph(text, style=style)
    return p


# ───────────────────────── Writer 组文档 ─────────────────────────

def build_doc_a():
    doc = Document()
    doc.add_heading("察元AI产品白皮书（测试稿）", level=0)
    para(doc, "察元AI文档助手是一款面向政企与个人的 WPS 智能加载项，2025年立项，"
              "目标是在编辑器内完成 AI 对话、文档审查与正文写回。本段用于改写测试："
              "它描述了产品从立项到落地的关键路径，包括需求调研、原型验证与首批客户试用三个阶段。")
    doc.add_heading("第一章 产品定位", level=1)
    para(doc, "察元AI定位为编辑器内的文档智能体，强调离线可用、内网可部署、写回安全三要素。")
    para(doc, "目标用户是经常与公文、报告、表格打交道的办公人群，以及需要批量校对文档的审校团队。")
    doc.add_heading("第二章 核心能力", level=1)
    para(doc, "核心能力包括：多文档错别字校对与批量修正、知识库检索问答、表格与演示稿的智能生成。")
    para(doc, "这段中文用于翻译测试。察元AI助手支持在段落后面直接插入译文，保持原文与译文对照阅读。")
    doc.add_heading("第三章 生态与接入", level=1)
    para(doc, "察元AI提供本机 MCP 文档智能体服务，外部智能体可以直读写本机 WPS 文档。")
    doc.save(os.path.join(OUT, "doc-A-内容.docx"))


def build_doc_b():
    doc = Document()
    doc.add_heading("季度运营报告（校对种子稿）", level=0)
    para(doc, "本季度销售额完成了年度目标的百分之三拾，用户增长稳订，其中新客户新增128家。")
    para(doc, "TODO：补充华东区明细数据。Lorem ipsum dolor sit amet，占位符文本需要清理。")
    para(doc, "用户可以通过登陆页面进入系统，登录后系统会自动同步配置；此前登陆逻辑存在偶发超时。")
    para(doc, "客服团队共受理工单1,024件，其中帐户类问题占比最高，账户安全巡检已按排。")
    para(doc, "备注：以上数据仅供参考，实际请以财务口径为准，请勿直接引用。")
    doc.save(os.path.join(OUT, "doc-B-审校.docx"))


def build_doc_c():
    doc = Document()
    doc.add_heading("项目背景", level=0)
    para(doc, "这是一个用于格式测试的普通段落，其中包含一个重要的关键词：重要事项需要高亮显示。")
    para(doc, "第二个段落用于对齐与行距测试。它需要足够长，以便在设置行距后仍能保持良好的可读性。")
    para(doc, "第三段用于缩进测试。首行缩进是中文公文排版的基本要求，通常为两个字符。")
    doc.save(os.path.join(OUT, "doc-C-格式.docx"))


def _add_bookmark(paragraph, name, bid):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def build_doc_d():
    doc = Document()
    doc.add_heading("组织架构与流程手册", level=0)
    h1 = doc.add_heading("第一章 总则", level=1)
    para(doc, "本章说明手册的适用范围与基本原则，适用于全体正式员工与外包协作人员。")
    h2 = doc.add_heading("第二章 组织架构", level=1)
    p2 = para(doc, "公司下设产品部、研发部、市场部与交付部四个一级部门，各部门负责人直接向总经理汇报。")
    _add_bookmark(p2, "察元书签甲", 101)
    h3 = doc.add_heading("第三章 业务流程", level=1)
    para(doc, "业务流程包括需求受理、方案评审、开发实施、验收交付四个环节，每个环节都有明确的责任人。")
    p3 = para(doc, "重大项目的流程变更必须经过变更委员会评审，评审通过后方可执行。")
    _add_bookmark(p3, "察元书签乙", 102)
    h4 = doc.add_heading("第四章 附则", level=1)
    para(doc, "本手册自发布之日起施行，由综合管理部负责解释与修订。")
    doc.save(os.path.join(OUT, "doc-D-结构.docx"))


def build_doc_e():
    doc = Document()
    doc.add_heading("图文混排测试文档", level=0)
    para(doc, "本文档包含一张预置图片与一个表格，用于图像与题注域测试。")
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(PNG_PATH, width=Pt(120))
    t = doc.add_table(rows=2, cols=3)
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "指标", "目标", "实际"
    row = t.rows[1].cells
    row[0].text, row[1].text, row[2].text = "营收", "1.0亿", "1.2亿"
    para(doc, "更多信息请访问察元官网了解详情。")
    doc.save(os.path.join(OUT, "doc-E-图文.docx"))


def build_doc_f():
    doc = Document()
    doc.add_heading("页面元素测试文档", level=0)
    for i in range(1, 9):
        para(doc, f"第{i}段：本段用于页眉页脚、水印与页面布局测试。"
                  f"段落内容填充：察元AI文档助手支持批量、确定性的文档写回操作{i}。")
    doc.save(os.path.join(OUT, "doc-F-页面.docx"))


def build_doc_g():
    doc = Document()
    doc.add_heading("脱敏测试文档", level=0)
    para(doc, "项目负责人张辉，联系电话13812345678，身份证号110101199003077758。")
    para(doc, "财务对接人李娜，银行卡号6222020200112233445，邮箱lina@example.com。")
    para(doc, "客户王强，手机号15900990088，家住北京市海淀区中关村大街1号。")
    doc.save(os.path.join(OUT, "doc-G-脱敏.docx"))


# ───────────────────────── 表格工作簿 ─────────────────────────

PRODUCTS = ["无人机", "火炮", "装甲车", "步枪", "导弹"]
REGIONS = ["华北", "华东", "华南", "西南"]
CUSTOMERS = ["前卫集团", "昆仑装备", "长城工业", "东海防务", "蓝鲸科技"]


def seed_sheet(ws, seed):
    rng = random.Random(seed)
    ws.append(["订单ID", "日期", "产品", "数量", "单价", "金额", "客户", "区域"])
    for i in range(1, 21):
        qty = rng.randint(1, 50)
        price = rng.randint(2, 90) * 1000
        ws.append([
            f"PO{seed:02d}{i:03d}",
            f"2026-{rng.randint(1, 9):02d}-{rng.randint(1, 28):02d}",
            rng.choice(PRODUCTS), qty, price, qty * price,
            rng.choice(CUSTOMERS), rng.choice(REGIONS),
        ])


def build_et():
    wb = Workbook()
    wb.remove(wb.active)
    for s in range(1, 61):
        ws = wb.create_sheet(f"S{s:02d}")
        seed_sheet(ws, s)
    wb.save(os.path.join(OUT, "e2e-et.xlsx"))


# ───────────────────────── 演示种子 ─────────────────────────

def build_wpp():
    prs = Presentation()
    tpl = prs.slide_layouts[0]
    sl = prs.slides.add_slide(tpl)
    sl.shapes.title.text = "察元AI年度汇报"
    sl.placeholders[1].text = "2026 年度经营回顾（种子页）"
    for title, body in [
        ("业务概览", "本页为种子页：业务概览，客户数与续约率稳步提升。"),
        ("数据亮点", "本页为种子页：数据亮点，营收完成目标120%。"),
        ("产品规划", "本页为种子页：产品规划，明年聚焦海外市场与产品线扩张。"),
        ("团队与协作", "本页为种子页：团队与协作，团队扩编至40人。"),
    ]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    end = prs.slides.add_slide(prs.slide_layouts[0])
    end.shapes.title.text = "谢谢观看"
    end.placeholders[1].text = "种子页：结束"
    prs.save(os.path.join(OUT, "e2e-wpp.pptx"))


if __name__ == "__main__":
    build_doc_a()
    build_doc_b()
    build_doc_c()
    build_doc_d()
    build_doc_e()
    build_doc_f()
    build_doc_g()
    build_et()
    build_wpp()
    print("seed files ->", os.path.abspath(OUT))
    for f in sorted(os.listdir(OUT)):
        print(" -", f, os.path.getsize(os.path.join(OUT, f)))
